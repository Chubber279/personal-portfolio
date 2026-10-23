"""Resume (1 page) and CV (2 pages) as .docx and Markdown, from content.py."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import content as c

FONT = "Calibri"
ACCENT_HEX = "222222"  # charcoal, matching the site's monochrome palette
ACCENT = RGBColor.from_string(ACCENT_HEX)

CONTACT_TEXT = [f"{c.LOCATION} · Open to remote", c.PHONE, c.EMAIL]
WORK_AUTH = " · ".join(c.WORK_AUTH)


def skill_line(items: list[str]) -> str:
    return ", ".join(items)


# --------------------------------------------------------------------------- docx helpers
def _set_font(run, size, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _spacing(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", ACCENT_HEX)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    pPr.append(borders)


def _hyperlink(p, text, url, size):
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = p.add_run(text)
    _set_font(run, size)
    run.font.underline = False
    link.append(run._r)
    p._p.append(link)


class Doc:
    def __init__(self, body_pt: float, margin_in: float):
        self.d = Document()
        self.body = body_pt
        s = self.d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(margin_in)
        s.top_margin = s.bottom_margin = Inches(margin_in - 0.1)
        self.width = 8.5 - 2 * margin_in
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(body_pt)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    def para(self, align=None):
        p = self.d.add_paragraph()
        _spacing(p)
        if align:
            p.alignment = align
        return p

    def run(self, p, text, size=None, bold=False, italic=False, color=None):
        r = p.add_run(text)
        _set_font(r, size or self.body, bold, italic, color)
        return r

    def header(self):
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER)
        self.run(p, c.NAME.upper(), 20, True, color=ACCENT)
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER)
        _spacing(p, 0, 2)
        self.run(p, c.TITLE, self.body + 1.5, color=ACCENT)
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER)
        self.run(p, "  |  ".join(CONTACT_TEXT))
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER)
        _spacing(p, 0, 4)
        for i, (text, url) in enumerate(c.LINKS):
            if i:
                self.run(p, "  |  ")
            _hyperlink(p, text, url, self.body)
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER)
        _spacing(p, 0, 4)
        self.run(p, WORK_AUTH, self.body - 0.5, italic=True)

    def section(self, title):
        p = self.para()
        _spacing(p, 6, 2)
        self.run(p, title.upper(), self.body + 0.5, True, color=ACCENT)
        _bottom_border(p)
        p.paragraph_format.keep_with_next = True

    def line_lr(self, left_runs, right_text, right_italic=True, after=0, before=0):
        """Left text with a right-aligned tab stop for dates."""
        p = self.para()
        _spacing(p, before, after)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(self.width), WD_TAB_ALIGNMENT.RIGHT)
        for text, bold, italic in left_runs:
            self.run(p, text, bold=bold, italic=italic)
        self.run(p, "\t" + right_text, italic=right_italic)
        p.paragraph_format.keep_with_next = True
        return p

    def bullet(self, text, after=1.5):
        p = self.d.add_paragraph(style="List Bullet")
        _spacing(p, 0, after)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.17)
        self.run(p, text)
        return p

    def text(self, text, after=0, italic=False, bold=False):
        p = self.para()
        _spacing(p, 0, after)
        self.run(p, text, italic=italic, bold=bold)
        return p

    def labelled(self, label, text, after=1):
        p = self.para()
        _spacing(p, 0, after)
        self.run(p, label + ": ", bold=True)
        self.run(p, text)
        return p

    def project(self, pr):
        self.line_lr([(pr["name"], True, False), (" | " + ", ".join(pr["tech"]), False, True)], "",
                     after=1)
        self.bullet(pr["text"])

    def save(self, path: Path):
        self.d.save(path)


# --------------------------------------------------------------------------- documents
def build_resume_docx(path: Path):
    d = Doc(body_pt=10, margin_in=0.55)
    d.header()
    d.section("Summary")
    d.text(c.SUMMARY)
    d.section("Skills")
    for label, items in c.SKILLS:
        d.labelled(label, skill_line(items), after=0.5)
    d.section("Experience")
    for i, job in enumerate(c.EXPERIENCE):
        d.line_lr([(job["company"], True, False), (" – " + job["role"], False, False)],
                  job["dates"], after=1, before=0 if i == 0 else 4)
        for b in job["resume"]:
            d.bullet(b)
    d.section("Projects")
    for pr in c.PROJECTS:
        d.project(pr)
    d.section("Education")
    ed = c.EDUCATION
    d.line_lr([(ed["school"], True, False), (" – " + ed["degree"], False, False)], ed["dates"], after=1)
    d.text(ed["detail"])
    d.section("Publication")
    pub = c.PUBLICATION
    d.text(pub["title"], bold=True)
    p = d.text(f"Ocampo & Clariño, {pub['short_venue']} (IEEE). Primary author; presented orally. DOI: ")
    _hyperlink(p, pub["doi"], pub["doi_url"], d.body)
    d.section("Certifications & Languages")
    for name, issuer, when in c.CERTIFICATIONS:
        d.line_lr([(name, True, False), (" – " + issuer, False, False)], when)
    d.labelled("Languages", "; ".join(f"{a} ({b})" for a, b in c.LANGUAGES), after=0)
    d.save(path)


def build_cv_docx(path: Path):
    d = Doc(body_pt=10.5, margin_in=0.75)
    d.header()
    d.section("Professional Summary")
    d.text(c.SUMMARY)
    d.section("Education")
    ed = c.EDUCATION
    d.line_lr([(ed["school"], True, False)], ed["dates"])
    d.text(f"{ed['degree']} · {ed['detail']}")
    d.section("Professional Experience")
    for i, job in enumerate(c.EXPERIENCE):
        d.line_lr([(job["company"], True, False)], job["dates"], before=0 if i == 0 else 5)
        d.text(f"{job['role']} · {job['location']}", after=1, italic=True)
        for b in job["cv"]:
            d.bullet(b)
    d.section("Projects")
    for pr in c.PROJECTS:
        d.project(pr)
    d.section("Research & Publications")
    pub = c.PUBLICATION
    d.text(pub["title"], bold=True)
    d.text(pub["authors"])
    p = d.text(pub["venue"] + " · 2024 · DOI: ")
    _hyperlink(p, pub["doi"], pub["doi_url"], d.body)
    d.bullet(pub["cv_line"])
    d.section("Conference Presentations")
    d.text(f"{pub['short_venue']}, IEEE", bold=True)
    d.bullet(c.PRESENTATION)
    d.section("Technical Skills")
    for label, items in c.SKILLS:
        d.labelled(label, skill_line(items))
    d.section("Certifications")
    for name, issuer, when in c.CERTIFICATIONS:
        d.line_lr([(name, True, False), (" – " + issuer, False, False)], when)
    d.section("Awards & Recognition")
    for a in c.AWARDS:
        d.bullet(a)
    d.section("Languages")
    for a, b in c.LANGUAGES:
        d.labelled(a, b, after=0)
    d.save(path)


# --------------------------------------------------------------------------- markdown
def build_md(kind: str) -> str:
    key = "resume" if kind == "resume" else "cv"
    ed, pub = c.EDUCATION, c.PUBLICATION
    out = [
        f"# {'Resume' if kind == 'resume' else 'Curriculum Vitae'} — Remote / Overseas",
        "",
        "<!--",
        "Generated by src/build.py from src/content.py — edit content.py, then re-run. Do not hand-edit.",
        "Purpose: submission-ready version tailored to remote, overseas, product-focused companies.",
        "Master source of truth for full detail remains the knowledge base (resume/Resume.md, resume/CV.md).",
        "-->",
        "",
        f"**{c.NAME}** — {c.TITLE}",
        "",
        " · ".join(CONTACT_TEXT) + " · " + " · ".join(f"[{t}]({u})" for t, u in c.LINKS),
        "",
        f"*{WORK_AUTH}*",
        "",
        "## Summary", "", c.SUMMARY, "",
    ]
    if kind == "cv":
        out += ["## Education", "", f"**{ed['school']}** · {ed['dates']}", "",
                f"{ed['degree']} · {ed['detail']}", ""]
    out += ["## Skills", ""] + [f"- **{a}:** {skill_line(b)}" for a, b in c.SKILLS] + [""]
    out += ["## Experience", ""]
    for job in c.EXPERIENCE:
        out += [f"### {job['company']} — {job['role']}", f"*{job['dates']} · {job['location']}*", ""]
        out += [f"- {b}" for b in job[key]] + [""]
    out += ["## Projects", ""]
    for pr in c.PROJECTS:
        out += [f"### {pr['name']} | {', '.join(pr['tech'])}", "", f"- {pr['text']}", ""]
    if kind == "resume":
        out += ["## Education", "", f"**{ed['school']}** — {ed['degree']} · {ed['dates']}", "",
                ed["detail"], ""]
    out += ["## Publication" if kind == "resume" else "## Research & Publications", "",
            f"**{pub['title']}**. {pub['authors']} {pub['venue']}. DOI: [{pub['doi']}]({pub['doi_url']})", "",
            pub[f"{key}_line"], ""]
    if kind == "cv":
        out += ["## Conference Presentations", "", f"- **{pub['short_venue']}, IEEE** — {c.PRESENTATION}", ""]
    out += ["## Certifications", ""] + [f"- **{n}** — {i}, {w}" for n, i, w in c.CERTIFICATIONS] + [""]
    if kind == "cv":
        out += ["## Awards & Recognition", ""] + [f"- {a}" for a in c.AWARDS] + [""]
    out += ["## Languages", ""] + [f"- **{a}:** {b}" for a, b in c.LANGUAGES] + [""]
    return "\n".join(out)


def build(out: Path, pdf: bool) -> list[Path]:
    """Write Markdown, .docx and (optionally) .pdf for both documents into `out`."""
    out.mkdir(parents=True, exist_ok=True)
    written = []
    for stem, kind, fn in ((c.RESUME_STEM, "resume", build_resume_docx), (c.CV_STEM, "cv", build_cv_docx)):
        md = out / f"{stem}.md"
        md.write_text(build_md(kind), encoding="utf-8")
        docx = out / f"{stem}.docx"
        fn(docx)
        written += [md, docx]
        if pdf:
            from docx2pdf import convert  # needs Microsoft Word

            convert(str(docx), str(out / f"{stem}.pdf"))
            written.append(out / f"{stem}.pdf")
    return written
